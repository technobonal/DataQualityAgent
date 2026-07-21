"""
rag/ingest.py

Ingesta del diccionario de datos: toma un documento que describe las
columnas del dataset (Excel/CSV, PDF, Word o texto plano) y lo convierte
en una lista de "chunks" (fragmentos de texto) listos para embeber y
guardar en el vector store (ver rag/vectorstore.py).

Dos caminos de procesamiento, según el formato:
  - .xlsx / .csv (estructurado): se espera una fila por columna del
    dataset, con campos como nombre/descripción/tipo/valores válidos.
    Cada fila se convierte en SU PROPIO chunk — mejor granularidad para
    la búsqueda: preguntar "qué significa la columna X" recupera
    exactamente esa definición, sin ruido de las demás columnas.
  - .pdf / .docx / .txt / .md (texto libre): se extrae todo el texto y
    se divide en chunks de tamaño fijo con superposición (overlap),
    como es estándar en pipelines de RAG para no cortar una idea justo
    en el límite entre dos fragmentos.

Sin dependencia de LLM — extracción y chunking son 100% deterministas.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field, asdict
from pathlib import Path
from typing import Any

import pandas as pd

_CHUNK_SIZE = 500     # caracteres por chunk (solo aplica a texto libre)
_CHUNK_OVERLAP = 50   # caracteres de superposición entre chunks consecutivos

_EXTENSIONES_ESTRUCTURADAS = {".xlsx", ".xls", ".csv"}
_EXTENSIONES_TEXTO_LIBRE = {".pdf", ".docx", ".txt", ".md"}


@dataclass
class DocumentChunk:
    """Un fragmento de texto listo para embeber, con su metadata de origen."""
    text: str
    source_file: str
    chunk_type: str  # 'structured_row' | 'free_text'
    metadata: dict[str, Any] = field(default_factory=dict)

    def to_dict(self) -> dict[str, Any]:
        """Serializa a dict plano — lo que rag/vectorstore.py guarda como metadata."""
        return asdict(self)


def _extraer_texto_pdf(path: Path) -> str:
    """Extrae todo el texto de un PDF, página por página."""
    import pypdf
    lector = pypdf.PdfReader(str(path))
    paginas = [pagina.extract_text() or "" for pagina in lector.pages]
    return "\n".join(paginas)


def _extraer_texto_docx(path: Path) -> str:
    """Extrae todo el texto de un Word: párrafos + contenido de tablas."""
    import docx
    documento = docx.Document(str(path))
    partes = [p.text for p in documento.paragraphs if p.text.strip()]

    for tabla in documento.tables:
        for fila in tabla.rows:
            texto_fila = " | ".join(celda.text.strip() for celda in fila.cells)
            if texto_fila.strip(" |"):
                partes.append(texto_fila)

    return "\n".join(partes)


def _dividir_en_chunks(texto: str, chunk_size: int = _CHUNK_SIZE, overlap: int = _CHUNK_OVERLAP) -> list[str]:
    """
    Divide un texto largo en fragmentos de tamaño fijo, con superposición
    entre fragmentos consecutivos — evita cortar una idea justo en el
    límite entre dos chunks.
    """
    texto = re.sub(r"\n{3,}", "\n\n", texto).strip()
    if not texto:
        return []

    chunks = []
    inicio = 0
    while inicio < len(texto):
        fin = inicio + chunk_size
        chunks.append(texto[inicio:fin])
        inicio += chunk_size - overlap

    return chunks


def _procesar_archivo_estructurado(path: Path) -> list[DocumentChunk]:
    """
    Procesa un diccionario de datos tabular (Excel/CSV): una fila por
    columna del dataset. Cada fila se convierte en su propio chunk.
    """
    df = pd.read_csv(path) if path.suffix.lower() == ".csv" else pd.read_excel(path)

    chunks = []
    for idx, fila in df.iterrows():
        partes_texto = [f"{col}: {valor}" for col, valor in fila.items() if pd.notna(valor)]
        texto_chunk = " | ".join(partes_texto)

        if not texto_chunk.strip():
            continue

        chunks.append(DocumentChunk(
            text=texto_chunk,
            source_file=path.name,
            chunk_type="structured_row",
            metadata={"row_index": int(idx), "columns": list(df.columns)},
        ))

    return chunks


def _procesar_archivo_texto_libre(path: Path) -> list[DocumentChunk]:
    """Procesa un diccionario de datos en texto libre (PDF/Word/TXT/MD)."""
    extension = path.suffix.lower()

    if extension == ".pdf":
        texto_completo = _extraer_texto_pdf(path)
    elif extension == ".docx":
        texto_completo = _extraer_texto_docx(path)
    else:  # .txt, .md
        texto_completo = path.read_text(encoding="utf-8", errors="ignore")

    fragmentos = _dividir_en_chunks(texto_completo)

    return [
        DocumentChunk(
            text=fragmento,
            source_file=path.name,
            chunk_type="free_text",
            metadata={"chunk_index": i, "total_chunks": len(fragmentos)},
        )
        for i, fragmento in enumerate(fragmentos)
    ]


def ingest_document(file_path: str) -> list[DocumentChunk]:
    """
    Punto de entrada principal: toma la ruta a un diccionario de datos
    (Excel, CSV, PDF o Word) y devuelve la lista de chunks listos para
    embeber y guardar en rag/vectorstore.py.

    Lanza FileNotFoundError si el archivo no existe, o ValueError si la
    extensión no está soportada.
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"No se encontró el archivo: {file_path}")

    extension = path.suffix.lower()

    if extension in _EXTENSIONES_ESTRUCTURADAS:
        return _procesar_archivo_estructurado(path)
    if extension in _EXTENSIONES_TEXTO_LIBRE:
        return _procesar_archivo_texto_libre(path)

    extensiones_validas = sorted(_EXTENSIONES_ESTRUCTURADAS | _EXTENSIONES_TEXTO_LIBRE)
    raise ValueError(f"Extensión no soportada: {extension!r}. Usar una de: {extensiones_validas}")


if __name__ == "__main__":
    import tempfile

    print("--- 1. Diccionario de datos estructurado (CSV) ---")
    with tempfile.NamedTemporaryFile(mode="w", suffix=".csv", delete=False, encoding="utf-8") as f:
        f.write("columna,descripcion,tipo,valores_validos\n")
        f.write("edad,Edad del cliente en años,numerico,0-120\n")
        f.write("estado,Estado de la cuenta,categorico,\"activo, inactivo, pendiente\"\n")
        ruta_csv = f.name

    chunks_csv = ingest_document(ruta_csv)
    for c in chunks_csv:
        print(" ", c.to_dict())

    print("\n--- 2. Diccionario de datos en texto libre (TXT) ---")
    with tempfile.NamedTemporaryFile(mode="w", suffix=".txt", delete=False, encoding="utf-8") as f:
        f.write(
            "Diccionario de datos del dataset de clientes.\n\n"
            "La columna edad representa la edad del cliente en años, "
            "y debe estar entre 0 y 120. " * 10 +
            "\n\nLa columna estado indica si la cuenta está activa, inactiva o pendiente."
        )
        ruta_txt = f.name

    chunks_txt = ingest_document(ruta_txt)
    print(f"  Se generaron {len(chunks_txt)} chunks")
    for c in chunks_txt[:2]:
        print(" ", c.to_dict())